"""Dựng báo cáo .docx đã sửa: python report/build_report.py  ->  C:/Users/Admin/Downloads/DATN-PTIT_sua.docx"""

import copy
import os
import sys

from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from docx_lib import Report, set_font
import content_ch1, content_ch2a, content_ch2b

ORIG = os.environ.get("ORIG", "C:/Users/Admin/Downloads/DATN-PTIT.docx")
OUT = os.environ.get("OUT", "C:/Users/Admin/Downloads/DATN-PTIT_sua.docx")
COVER_LAST = 16     # chỉ số đoạn cuối của trang bìa (P0016)


def make(labels=None):
    R = Report(ORIG)
    if labels:
        R.labels = dict(labels)
    body = R.doc.element.body
    keep = R.doc.paragraphs[COVER_LAST]._p
    after = False
    for el in list(body):
        if el is keep:
            after = True
            continue
        if after and el.tag != qn("w:sectPr"):
            body.remove(el)
    sect = body.find(qn("w:sectPr"))
    # đoạn cuối trang bìa đã chứa ngắt trang
    content_ch1.build(R, ORIG)
    R.page_break()
    content_ch2a.build(R)
    content_ch2b.build(R)
    R.page_break()
    content_ch2b.appendix_and_refs(R)
    # số trang ở chân trang
    fp = R.doc.sections[0].footer.paragraphs[0] if R.doc.sections[0].footer.paragraphs else R.doc.sections[0].footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if not fp.text.strip():
        r = fp.add_run()
        set_font(r, 11)
        for t, txt in (("begin", None), ("instr", " PAGE "), ("end", None)):
            if t == "instr":
                e = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve">{txt}</w:instrText>')
            else:
                e = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="{t}"/>')
            r._r.append(e)
    return R


if __name__ == "__main__":
    R1 = make()
    R2 = make(R1.labels)
    assert R1.labels == R2.labels
    R2.doc.save(OUT)
    print("saved", OUT, "| công thức:", sum(R2.eq_no.values()), "hình:", sum(R2.fig_no.values()), "bảng:", sum(R2.tab_no.values()))
